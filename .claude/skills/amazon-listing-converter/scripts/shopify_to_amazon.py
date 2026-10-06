#!/usr/bin/env python3
"""shopify_to_amazon - convert Foxy Printing (Shopify) products into a safe Amazon UK upload pack.

  python3 shopify_to_amazon.py INPUT [INPUT ...] --out OUTDIR [--config amazon_config.json]
                               [--template Amazon-category-template.xlsm] [--only-personalised]

INPUT is either a Shopify product CSV export (Admin > Products > Export) or JSON from the Shopify
Admin GraphQL API (a list of product nodes, or {"data":{"products":{"nodes":[...]}}}, or JSON lines).

Writes to OUTDIR:
  Amazon upload - <date>.xlsx          review workbook: Listings, Amazon Custom, Custom templates,
                                       Held back, Checks, How to upload
  Amazon JSON_LISTINGS_FEED - <date>.json   the SP-API feed (Selling Partner connector / API)
  Amazon Custom mapping - <date>.txt   SKU -> customisation template, tab-delimited
  Amazon template FILLED - <name>      only with --template: Amazon's own category template, filled in
  summary.json

Safety: nothing is sent to Amazon. Products are HELD BACK (not converted) when they have no SKU,
duplicate SKUs, no price, no image, no product-type mapping, or possible trademark/celebrity content.
"""
import argparse, copy, csv, datetime, html, json, os, re, sys
from collections import OrderedDict

HERE = os.path.dirname(os.path.abspath(__file__))


# ----------------------------------------------------------------------------- input
def load_products(paths):
    prods = []
    for p in paths:
        if p.lower().endswith(".csv"):
            prods += from_shopify_csv(p)
        else:
            raw = open(p, encoding="utf-8").read().strip()
            try:
                data = json.loads(raw)
                items = find_nodes(data)
            except json.JSONDecodeError:  # JSON lines
                items = []
                for line in raw.splitlines():
                    if line.strip():
                        items += find_nodes(json.loads(line))
            prods += [from_graphql(n) for n in items]
    return prods


def find_nodes(d):
    if isinstance(d, list):
        out = []
        for x in d:
            out += find_nodes(x) if not (isinstance(x, dict) and "variants" in x) else [x]
        return out
    if isinstance(d, dict):
        if "variants" in d and ("title" in d or "handle" in d):
            return [d]
        for k in ("data", "products", "nodes", "edges", "node"):
            if k in d:
                return find_nodes(d[k])
    return []


def _nodes(x):
    if x is None:
        return []
    if isinstance(x, list):
        return [e.get("node", e) for e in x]
    return [e.get("node", e) for e in (x.get("nodes") or x.get("edges") or [])]


def from_graphql(n):
    imgs = []
    for m in _nodes(n.get("media")) + _nodes(n.get("images")):
        url = (m.get("image") or {}).get("url") or m.get("url") or m.get("src")
        if url and url not in imgs:
            imgs.append(url)
    variants = []
    for v in _nodes(n.get("variants")):
        vimg = None
        for m in _nodes(v.get("media")):
            vimg = (m.get("image") or {}).get("url") or vimg
        if v.get("image"):
            vimg = v["image"].get("url") or vimg
        variants.append({"sku": (v.get("sku") or "").strip(), "price": v.get("price"),
                         "barcode": (v.get("barcode") or "").strip(),
                         "options": OrderedDict((o["name"], o["value"]) for o in v.get("selectedOptions", [])
                                                if o["name"] != "Title"),
                         "image": vimg})
    return {"id": n.get("id", ""), "handle": n.get("handle", ""), "title": n.get("title", ""),
            "status": (n.get("status") or "ACTIVE").upper(), "vendor": n.get("vendor", ""),
            "product_type": n.get("productType", ""), "tags": n.get("tags", []),
            "description_html": n.get("descriptionHtml", ""),
            "options": [o["name"] for o in n.get("options", []) if o["name"] != "Title"],
            "images": imgs, "variants": variants}


def from_shopify_csv(path):
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    by = OrderedDict()
    for r in rows:
        h = r.get("Handle", "").strip()
        if not h:
            continue
        p = by.get(h)
        if p is None:
            p = by[h] = {"id": h, "handle": h, "title": r.get("Title", ""), "status": (r.get("Status") or "active").upper(),
                         "vendor": r.get("Vendor", ""), "product_type": r.get("Type", ""),
                         "tags": [t.strip() for t in r.get("Tags", "").split(",") if t.strip()],
                         "description_html": r.get("Body (HTML)", ""),
                         "options": [r[f"Option{i} Name"] for i in (1, 2, 3)
                                     if r.get(f"Option{i} Name") and r.get(f"Option{i} Name") != "Title"],
                         "images": [], "variants": []}
        src = (r.get("Image Src") or "").strip()
        if src and src not in p["images"]:
            p["images"].append(src)
        if (r.get("Variant SKU") or r.get("Variant Price") or "").strip():
            opts = OrderedDict()
            for i, name in enumerate(p["options"], 1):
                val = r.get(f"Option{i} Value", "")
                if val:
                    opts[name] = val
            p["variants"].append({"sku": r.get("Variant SKU", "").strip(), "price": r.get("Variant Price"),
                                  "barcode": (r.get("Variant Barcode") or "").strip().lstrip("'"),
                                  "options": opts, "image": (r.get("Variant Image") or "").strip() or None})
    return list(by.values())


# ----------------------------------------------------------------------------- text
def clean_text(s, cfg):
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s or "")
    s = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>|</div>", "\n", s)
    s = re.sub(r"(?i)<li[^>]*>", "\n• ", s)
    s = html.unescape(re.sub(r"<[^>]+>", "", s))
    s = re.sub(r"(?i)[\w.+-]+@[\w-]+(\.[\w-]+)+", "", s)                          # emails
    s = re.sub(r"(?i)\b(https?://\S+|www\.\S+|[\w-]+\.(co\.uk|com|net|org)\S*)", "", s)  # urls/domains
    s = re.sub(r"(?<!\d)(\+44\s?|0)\d{3,4}[\s-]?\d{3}[\s-]?\d{3,4}(?!\d)", "", s)   # UK phone numbers
    for rx in cfg["strip_text_regex"]:
        s = re.sub(rx, "", s)
    s = re.sub(r"[ \t]+", " ", s)
    lines = [l.strip(" -–•\t") for l in s.splitlines()]
    return [l for l in lines if len(re.findall(r"[A-Za-z]", l)) >= 3]


def join_words(items):
    items = [i.lower() for i in items]
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def li_items(html_s):
    return [html.unescape(re.sub(r"<[^>]+>", "", m)).strip()
            for m in re.findall(r"(?is)<li[^>]*>(.*?)</li>", html_s or "")]


def clean_title(t, cfg):
    t = html.unescape(t)
    t = "".join(ch for ch in t if ch not in cfg["banned_title_chars"])
    return re.sub(r"\s+", " ", t).strip()[:200]


def cut_bytes(s, n):
    while len(s.encode()) > n:
        s = s.rsplit(" ", 1)[0] if " " in s else s[:-1]
    return s


# ----------------------------------------------------------------------------- rules
def product_type_for(p, cfg):
    for rule in cfg["product_types"]:
        if re.search(rule["match"], p["product_type"] or "") or re.search(rule["match"], p["title"]):
            return rule
    return None


def personalisation(p, cfg):
    fields, why = [], []
    tags = p["tags"]
    text = " ".join(clean_text(p["description_html"], dict(cfg, strip_text_regex=[])))
    for r in cfg["personalisation_rules"]:
        hit = (("tag" in r and any(t.lower() == r["tag"].lower() for t in tags)) or
               ("tag_regex" in r and any(re.search(r["tag_regex"], t) for t in tags)) or
               ("title_regex" in r and re.search(r["title_regex"], p["title"])) or
               ("text_regex" in r and re.search(r["text_regex"], text)))
        if hit:
            why.append(r.get("tag") or r.get("tag_regex") or r.get("title_regex") or r.get("text_regex"))
            for f in r["fields"]:
                if f not in fields:
                    fields.append(f)
    return fields, why


def field_spec(name, cfg):
    d = dict(cfg["personalisation_fields"].get("_default"))
    d.update(cfg["personalisation_fields"].get(name, {}))
    d.setdefault("required", True)
    return d


def ip_check(p, cfg, skus):
    if any(s in cfg.get("ip_allow_skus", []) for s in skus):
        return []
    hay = " ".join([p["title"], p["product_type"], " ".join(p["tags"]), p["description_html"]])
    hits = [t for t in cfg["ip_hold_terms"] if re.search(r"(?i)(?<![\w])" + re.escape(t) + r"(?![\w])", hay)]
    hits += [f"tag '{t}'" for t in p["tags"] if re.search(cfg["ip_hold_tags_regex"], t)]
    return hits


def price_for(v, cfg):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    rule = cfg.get("price_rule", "same")
    if rule.startswith("plus:"):
        x += float(rule[5:])
    elif rule.startswith("markup:"):
        x *= 1 + float(rule[7:]) / 100
    return round(x + 1e-9, 2)


def valid_ean(code):
    if not re.fullmatch(r"\d{8}|\d{12,14}", code or ""):
        return False
    digits = [int(c) for c in code]
    check = digits.pop()
    s = sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(digits)))
    return (10 - s % 10) % 10 == check


# ----------------------------------------------------------------------------- build
def keywords(p, title, cfg, ptype=None, personalised=False):
    words, seen = [], set(re.findall(r"[a-z0-9']+", title.lower()))
    rule_tags = {r["tag"].lower() for r in cfg["personalisation_rules"] if "tag" in r}
    extra = list((ptype or {}).get("keywords", []))
    if personalised:
        extra += cfg.get("personalised_keywords", [])
    for t in p["tags"] + extra:
        if (re.search(cfg["internal_tags_regex"], t) or re.search(cfg["ip_hold_tags_regex"], t)
                or t.lower() in rule_tags or re.match(r"(?i)^add\b", t)):
            continue
        for w in re.findall(r"[a-z0-9']+", t.lower().replace("-", " ")):
            if w not in seen and len(w) > 2 and w not in cfg.get("keyword_block", []):
                seen.add(w)
                words.append(w)
    return cut_bytes(" ".join(words), 249)


def bullets_for(p, ptype, fields, cfg):
    out = []
    for b in li_items(p["description_html"]):
        b = " ".join(clean_text(b, cfg))
        if 15 <= len(b) <= 500 and b not in out:
            out.append(b)
    if fields:
        out.insert(0, f"PERSONALISED: click 'Customise now' and enter the {join_words(fields)} - we print it exactly as typed")
    for b in ptype.get("defaults", {}).get("bullets", []):
        if len(out) >= 5:
            break
        if b not in out:
            out.append(b)
    return out[:5]


def description_for(p, fields, cfg):
    lines = clean_text(p["description_html"], cfg)
    lines = [l for l in lines if not l.startswith("•")] or lines
    txt = "\n".join(lines)
    if fields:
        txt += ("\n\nPersonalise it: use the Customise now button to enter the " + join_words(fields) +
                ". Please check spelling - we print exactly what you type.")
    return txt[:2000].strip()


def convert(products, cfg, only_personalised=False):
    listings, held, checks, custom_rows, templates = [], [], [], [], OrderedDict()
    seen_skus = {}
    for p in products:
        for v in p["variants"]:
            if v["sku"]:
                seen_skus[v["sku"]] = seen_skus.get(v["sku"], 0) + 1
    for p in products:
        skus = [v["sku"] for v in p["variants"]]
        reason = []
        if p["status"] not in cfg["include_statuses"]:
            reason.append(f"status {p['status']}")
        if not p["variants"] or any(not s for s in skus):
            reason.append("variant without SKU")
        dups = [s for s in skus if s and seen_skus.get(s, 0) > 1]
        if dups:
            reason.append("duplicate SKU " + ", ".join(sorted(set(dups))))
        if any(len(s) > 40 for s in skus):
            reason.append("SKU longer than 40 characters")
        if any(price_for(v["price"], cfg) in (None, 0) for v in p["variants"]):
            reason.append("missing price")
        if not p["images"]:
            reason.append("no images")
        ptype = product_type_for(p, cfg)
        if not ptype:
            reason.append(f"no Amazon product type mapped for '{p['product_type']}'")
        ip = ip_check(p, cfg, skus)
        if ip:
            reason.append("possible trademark / celebrity content: " + ", ".join(ip[:4]))
        fields, why = personalisation(p, cfg)
        if only_personalised and not fields:
            reason.append("not personalised (--only-personalised)")
        if reason:
            held.append({"handle": p["handle"], "title": p["title"], "skus": ", ".join(skus), "reason": "; ".join(reason)})
            continue
        title = clean_title(p["title"], cfg)
        warn = []
        for t in cfg.get("warn_terms", []):
            if re.search(r"(?i)\b" + re.escape(t) + r"\b", p["title"] + " " + " ".join(p["tags"])):
                warn.append(f"contains '{t}' - check Amazon content/IP policy")
        words = re.findall(r"[a-z']+", title.lower())
        rep = sorted({w for w in words if len(w) > 3 and words.count(w) > 2})
        if rep:
            warn.append("title repeats " + ", ".join(rep) + " more than twice (Amazon title rule)")
        if len(p["title"]) > 200:
            warn.append("title cut to 200 characters")
        if title.isupper():
            warn.append("title is ALL CAPS")
        d = ptype.get("defaults", {})
        lead = cfg["handling_days_personalised"] if fields else cfg["handling_days"]
        base = {
            "handle": p["handle"], "shopify_id": p["id"], "product_type": ptype["product_type"],
            "item_type_keyword": ptype.get("item_type_keyword", ""), "brand": cfg["brand"],
            "manufacturer": cfg["manufacturer"], "title": title,
            "description": description_for(p, fields, cfg), "bullets": bullets_for(p, ptype, fields, cfg),
            "keywords": keywords(p, title, cfg, ptype, bool(fields)), "material": d.get("material", ""), "color": d.get("color", ""),
            "capacity_ml": d.get("capacity_ml"), "included_components": d.get("included_components", ""),
            "images": p["images"][:9], "personalised": bool(fields), "fields": fields,
            "handling_days": lead, "warnings": warn,
        }
        tname = ""
        if fields:
            tname = "FOXY - " + " + ".join(fields)
            templates.setdefault(tname, [dict(field=f, **field_spec(f, cfg)) for f in fields])
        opts = p["options"]
        multi = len(p["variants"]) > 1 and opts
        if multi:
            maps = [cfg["variation_themes"].get(o.lower(), cfg["variation_themes"]["_default"]) for o in opts]
            theme = "/".join(m["theme"] for m in maps)
            attrs = [m["attribute"] for m in maps]
            if len(set(attrs)) < len(attrs):
                held.append({"handle": p["handle"], "title": p["title"], "skus": ", ".join(skus),
                             "reason": f"options {opts} map to the same Amazon attribute - add them to variation_themes"})
                continue
            stem = re.sub(r"[-_ ]\d*$", "", os.path.commonprefix(skus)).rstrip("-_ ") or skus[0]
            parent_sku = (stem + "-PARENT")[-40:] if len(stem) > 33 else stem + "-PARENT"
            if parent_sku in seen_skus:
                parent_sku = (parent_sku[:33] + "-PAR-AZ")[:40]
            listings.append(dict(base, sku=parent_sku, parentage="parent", parent_sku="", variation_theme=theme,
                                 price=None, quantity=None, ean="", variation_values={}))
            for v in p["variants"]:
                vals = {attrs[i]: v["options"].get(o, "") for i, o in enumerate(opts)}
                imgs = ([v["image"]] if v["image"] else []) + [u for u in p["images"] if u != v["image"]]
                listings.append(dict(base, sku=v["sku"], parentage="child", parent_sku=parent_sku,
                                     variation_theme=theme, price=price_for(v["price"], cfg), quantity=cfg["quantity"],
                                     ean=v["barcode"] if valid_ean(v["barcode"]) else "", variation_values=vals,
                                     title=clean_title(f"{title} - {' '.join(vals.values())}", cfg), images=imgs[:9]))
            if len(opts) > 1:
                base["warnings"].append(f"two-option theme {theme} - confirm it is allowed for {ptype['product_type']}")
        else:
            v = p["variants"][0]
            listings.append(dict(base, sku=v["sku"], parentage="", parent_sku="", variation_theme="",
                                 price=price_for(v["price"], cfg), quantity=cfg["quantity"],
                                 ean=v["barcode"] if valid_ean(v["barcode"]) else "", variation_values={}))
        for row in listings:
            if row["handle"] == p["handle"] and row["parentage"] != "parent":
                if fields:
                    custom_rows.append({"sku": row["sku"], "title": row["title"], "template": tname,
                                        "fields": ", ".join(fields)})
                for b in [x for x in p["variants"] if x["sku"] == row["sku"] and x["barcode"] and not valid_ean(x["barcode"])]:
                    row["warnings"].append(f"barcode {b['barcode']} is not a valid EAN - sent as GTIN-exempt")
        for w in base["warnings"]:
            checks.append({"sku": skus[0], "title": title, "check": w})
    return listings, held, checks, custom_rows, templates


# ----------------------------------------------------------------------------- outputs
def L(value, cfg, lang=True):
    d = {"value": value, "marketplace_id": cfg["marketplace_id"]}
    if lang:
        d["language_tag"] = cfg["language_tag"]
    return [d]


def json_attributes(r, cfg):
    mp = cfg["marketplace_id"]
    a = OrderedDict()
    a["item_name"] = L(r["title"], cfg)
    a["brand"] = L(r["brand"], cfg)
    a["manufacturer"] = L(r["manufacturer"], cfg)
    a["condition_type"] = [{"value": "new_new", "marketplace_id": mp}]
    if r["item_type_keyword"]:
        a["item_type_keyword"] = [{"value": r["item_type_keyword"], "marketplace_id": mp}]
    if r["parentage"] != "parent":
        if r["ean"]:
            a["externally_assigned_product_identifier"] = [{"type": "ean", "value": r["ean"], "marketplace_id": mp}]
        else:
            a["supplier_declared_has_product_identifier_exemption"] = [{"value": True, "marketplace_id": mp}]
    a["product_description"] = L(r["description"], cfg)
    a["bullet_point"] = [{"value": b, "language_tag": cfg["language_tag"], "marketplace_id": mp} for b in r["bullets"]]
    if r["keywords"]:
        a["generic_keyword"] = L(r["keywords"], cfg)
    if r["material"]:
        a["material"] = L(r["material"], cfg)
    if r["color"] and "color" not in r["variation_values"] and "COLOR" not in r["variation_theme"]:
        a["color"] = L(r["color"], cfg)
    if r["capacity_ml"]:
        a["capacity"] = [{"value": r["capacity_ml"], "unit": "milliliters", "marketplace_id": mp}]
    if r["included_components"]:
        a["included_components"] = L(r["included_components"], cfg)
    a["number_of_items"] = [{"value": 1, "marketplace_id": mp}]
    a["country_of_origin"] = [{"value": cfg["country_of_origin"], "marketplace_id": mp}]
    a["part_number"] = [{"value": r["sku"], "marketplace_id": mp}]
    for k, v in r["variation_values"].items():
        a[k] = L(v, cfg)
    if r["parentage"]:
        a["parentage_level"] = [{"value": r["parentage"], "marketplace_id": mp}]
        a["variation_theme"] = [{"name": r["variation_theme"]}]
    if r["parentage"] == "child":
        a["child_parent_sku_relationship"] = [{"child_relationship_type": "variation", "parent_sku": r["parent_sku"],
                                               "marketplace_id": mp}]
    if r["parentage"] != "parent":
        a["purchasable_offer"] = [{"currency": cfg["currency"], "marketplace_id": mp,
                                   "our_price": [{"schedule": [{"value_with_tax": r["price"]}]}]}]
        a["fulfillment_availability"] = [{"fulfillment_channel_code": "DEFAULT", "quantity": r["quantity"],
                                          "lead_time_to_ship_max_days": r["handling_days"]}]
    imgs = r["images"]
    if imgs:
        a["main_product_image_locator"] = [{"media_location": imgs[0], "marketplace_id": mp}]
        for i, u in enumerate(imgs[1:9], 1):
            a[f"other_product_image_locator_{i}"] = [{"media_location": u, "marketplace_id": mp}]
    return a


def write_json_feed(listings, cfg, path):
    msgs = []
    for i, r in enumerate(listings, 1):
        msgs.append({"messageId": i, "sku": r["sku"], "operationType": "UPDATE", "productType": r["product_type"],
                     "requirements": "LISTING", "attributes": json_attributes(r, cfg)})
    feed = {"header": {"sellerId": cfg.get("seller_id") or "YOUR_SELLER_ID", "version": "2.0",
                       "issueLocale": cfg["language_tag"]}, "messages": msgs}
    json.dump(feed, open(path, "w", encoding="utf-8"), indent=1, ensure_ascii=False)


LISTING_COLS = ["sku", "parentage", "parent_sku", "variation_theme", "variation", "product_type", "title", "brand",
                "price", "quantity", "handling_days", "ean / exemption", "personalised", "custom template",
                "bullet 1", "bullet 2", "bullet 3", "bullet 4", "bullet 5", "keywords", "description",
                "main image", "other images", "warnings", "shopify handle"]


def listing_row(r, tmpl):
    b = r["bullets"] + [""] * 5
    return [r["sku"], r["parentage"] or "standalone", r["parent_sku"], r["variation_theme"],
            ", ".join(f"{k}={v}" for k, v in r["variation_values"].items()), r["product_type"], r["title"], r["brand"],
            r["price"], r["quantity"], r["handling_days"],
            "" if r["parentage"] == "parent" else (r["ean"] or "GTIN exempt"),
            "YES" if r["personalised"] else "", tmpl, *b[:5], r["keywords"], r["description"],
            r["images"][0] if r["images"] else "", " | ".join(r["images"][1:]), "; ".join(r["warnings"]), r["handle"]]


HOW_TO = [
    "SAFE UPLOAD ORDER - nothing in this pack has been sent to Amazon.",
    "",
    "1. Review the Listings sheet (titles, prices, bullets) and the Held back sheet. Fix anything in Shopify and re-run.",
    "2. Confirm each product_type (DRINKING_CUP for mugs) in Seller Central > Catalogue > Add products via upload >",
    "   search the product type, and download that category template.",
    "3a. RECOMMENDED - Amazon's own template: re-run with --template <downloaded template> and upload the",
    "    'Amazon template FILLED' file in Add products via upload. Amazon checks it and shows a processing report.",
    "3b. OR with API access (Amazon Selling Partner connector / SP-API): submit each message of the JSON feed with",
    "    putListingsItem mode=VALIDATION_PREVIEW first, fix every issue, then submit for real (or send the whole",
    "    JSON_LISTINGS_FEED through the Feeds API). Fill header.sellerId first.",
    "4. Wait until the listings are live (each child SKU gets an ASIN).",
    "5. PERSONALISATION (Amazon Custom) - needed so customers get the name box:",
    "   a. Professional seller account, enrolled in Amazon Custom (Seller Central > Amazon Custom). Custom",
    "      items must be Fulfilled by You (no FBA) - this pack already sends fulfilment channel DEFAULT.",
    "   b. For each template on the 'Custom templates' sheet, open ONE SKU from that group > Edit >",
    "      Customisation information > add a Text input (surface) per field with the label, instructions and",
    "      character limit shown, fix the font/colour (customers only type the text), add the preview area.",
    "   c. Then apply that saved template to the other SKUs in bulk with Amazon's bulk customisation tool",
    "      (Amazon Custom > bulk customisations), using the 'Amazon Custom mapping' file for which SKU gets which template.",
    "   d. Buy-test one personalised listing yourself (or check the order report) - customisation text arrives",
    "      in the order's customisation file (Orders API BuyerCustomizedInfo).",
    "6. Re-run this converter whenever Shopify changes; SKUs match Shopify so stock/price tools can link them.",
]


def write_workbook(listings, held, checks, custom_rows, templates, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0D0D0D")
    orange = PatternFill("solid", fgColor="FFE3CC")

    def sheet(title, cols, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = (widths or {}).get(c, min(max(len(c) + 4, 14), 60))
        return ws

    wb.remove(wb.active)
    tmpl_by_sku = {c["sku"]: c["template"] for c in custom_rows}
    ws = sheet("Listings", LISTING_COLS, [listing_row(r, tmpl_by_sku.get(r["sku"], "")) for r in listings],
               {"title": 60, "description": 60, "keywords": 40, "warnings": 50})
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "parent":
            for c in row:
                c.fill = orange
    sheet("Amazon Custom", ["sku", "title", "template", "fields"],
          [[c["sku"], c["title"], c["template"], c["fields"]] for c in custom_rows], {"title": 70, "template": 34})
    trows = []
    for name, fs in templates.items():
        for f in fs:
            trows.append([name, f["field"], "Text", f["max_chars"], "Yes" if f["required"] else "No", f["instructions"],
                          "Fixed by design (customer types text only)"])
    sheet("Custom templates", ["template", "field label", "input type", "max characters", "required",
                               "customer instructions", "font / colour"], trows,
          {"template": 34, "customer instructions": 60, "font / colour": 38})
    sheet("Held back", ["handle", "title", "skus", "reason"],
          [[h["handle"], h["title"], h["skus"], h["reason"]] for h in held], {"title": 60, "reason": 80})
    sheet("Checks", ["sku", "title", "check"], [[c["sku"], c["title"], c["check"]] for c in checks],
          {"title": 60, "check": 80})
    ws = wb.create_sheet("How to upload")
    for line in HOW_TO:
        ws.append([line])
    ws.column_dimensions["A"].width = 130
    ws["A1"].font = Font(bold=True)
    wb.move_sheet("How to upload", offset=-len(wb.sheetnames) + 1)
    wb.save(path)


# ------------------------------------------------------------- fill Amazon's own category template
def template_value(key, r, cfg):
    """Map an Amazon template column key (old flat-file or new attribute style) to our value."""
    k = key.strip().lower().lstrip(":")
    base = re.sub(r"\[.*?\]", "", k)
    m = re.match(r"([a-z_]+?)(\d*)(?:#(\d+))?(?:\.(.*))?$", base)
    if not m:
        return None
    name, num, idx, sub = m.group(1).rstrip("_"), m.group(2), m.group(3), m.group(4) or ""
    n = int(num or idx or 1)
    par = r["parentage"]
    img = r["images"]
    if name in ("item_sku", "contribution_sku", "sku"):
        return r["sku"]
    if name in ("feed_product_type", "product_type"):
        return r["product_type"]
    if name == "item_name":
        return r["title"]
    if name in ("brand_name", "brand"):
        return r["brand"]
    if name == "manufacturer":
        return r["manufacturer"]
    if name == "product_description":
        return r["description"]
    if name == "bullet_point":
        return r["bullets"][n - 1] if n - 1 < len(r["bullets"]) else None
    if name in ("generic_keywords", "generic_keyword"):
        return r["keywords"]
    if name == "record_action":
        return "Create or Replace (Full Update)"
    if name in ("update_delete", "operation_type"):
        return "Update"
    if name == "item_type_keyword" or name == "item_type":
        return r["item_type_keyword"]
    if name in ("condition_type",):
        return "New"
    if name in ("standard_price", "our_price") or (name == "purchasable_offer" and "value_with_tax" in sub):
        return r["price"]
    if name == "quantity" or (name == "fulfillment_availability" and sub.endswith("quantity")):
        return r["quantity"]
    if name in ("fulfillment_latency",) or (name == "fulfillment_availability" and "lead_time" in sub):
        return r["handling_days"] if par != "parent" else None
    if name == "main_image_url" or (name == "main_product_image_locator"):
        return img[0] if img else None
    if name in ("other_image_url", "other_product_image_locator"):
        i = n if name == "other_image_url" else int(num or 1)
        return img[i] if i < len(img) else None
    if name in ("parent_child", "parentage_level"):
        return {"parent": "Parent", "child": "Child"}.get(par)
    if name == "parent_sku" or (name == "child_parent_sku_relationship" and "parent_sku" in sub):
        return r["parent_sku"] or None
    if name in ("relationship_type",) or (name == "child_parent_sku_relationship" and "relationship_type" in sub):
        return "Variation" if par == "child" else None
    if name == "variation_theme":
        return r["variation_theme"] or None
    if name in ("color_name", "color"):
        if "COLOR" in r["variation_theme"]:
            return r["variation_values"].get("color")
        return r["color"] or None
    if name in ("size_name", "size"):
        return r["variation_values"].get("size")
    if name in ("style_name", "style"):
        return r["variation_values"].get("style")
    if name in ("material_type", "material"):
        return r["material"] or None
    if name in ("external_product_id",) or (name == "externally_assigned_product_identifier" and "value" in sub):
        return r["ean"] or None
    if name in ("external_product_id_type",) or (name == "externally_assigned_product_identifier" and "type" in sub):
        return "EAN" if r["ean"] else None
    if name == "supplier_declared_has_product_identifier_exemption":
        return "Yes" if (par != "parent" and not r["ean"]) else None
    if name in ("country_of_origin",):
        return "United Kingdom"
    if name == "part_number":
        return r["sku"]
    if name == "number_of_items":
        return 1
    if name in ("capacity",) and "unit" in sub:
        return "Millilitres" if r["capacity_ml"] else None
    if name in ("capacity",):
        return r["capacity_ml"]
    if name in ("included_components",):
        return r["included_components"] or None
    return None


def fill_template(tpath, listings, cfg, out_path):
    from openpyxl import load_workbook
    wb = load_workbook(tpath, keep_vba=tpath.lower().endswith(".xlsm"))
    ws = next((wb[s] for s in wb.sheetnames if s.lower() == "template"), None) or \
        next((wb[s] for s in wb.sheetnames if "template" in s.lower()), wb.active)
    key_row, best = None, 0
    for r in range(1, 12):
        vals = [str(c.value or "") for c in ws[r]]
        score = sum(1 for v in vals if re.match(r"^[a-z_]+[\w#\[\]=.\-]*$", v))
        if score > best and any(re.search(r"(sku|item_name)", v) for v in vals):
            key_row, best = r, score
    if not key_row:
        sys.exit("could not find the attribute-name row in the template")
    keys = [str(c.value or "") for c in ws[key_row]]
    sku_col = next(i for i, k in enumerate(keys) if re.search(r"(item_sku|contribution_sku|^sku)", k))
    row = key_row + 1
    while ws.cell(row, sku_col + 1).value not in (None, ""):
        if str(ws.cell(row, sku_col + 1).value).lower().startswith(("abc", "example", "sku")):
            row += 1  # leave Amazon's example rows
            continue
        row += 1
    filled, unmapped = set(), set(keys) - {""}
    for r in listings:
        for ci, k in enumerate(keys):
            if not k:
                continue
            v = template_value(k, r, cfg)
            if v not in (None, ""):
                ws.cell(row, ci + 1, v)
                filled.add(k)
        row += 1
    wb.save(out_path)
    return sorted(filled), sorted(unmapped - filled)


# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--out", default="amazon-upload")
    ap.add_argument("--config", default=os.path.join(HERE, "amazon_config.json"))
    ap.add_argument("--template", help="Amazon category template (.xlsm/.xlsx) downloaded from Seller Central")
    ap.add_argument("--only-personalised", action="store_true")
    a = ap.parse_args()
    cfg = json.load(open(a.config, encoding="utf-8"))
    os.makedirs(a.out, exist_ok=True)
    today = datetime.date.today().isoformat()
    products = load_products(a.inputs)
    listings, held, checks, custom_rows, templates = convert(products, cfg, a.only_personalised)
    xl = os.path.join(a.out, f"Amazon upload - {today}.xlsx")
    feed = os.path.join(a.out, f"Amazon JSON_LISTINGS_FEED - {today}.json")
    mapping = os.path.join(a.out, f"Amazon Custom mapping - {today}.txt")
    write_workbook(listings, held, checks, custom_rows, templates, xl)
    write_json_feed(listings, cfg, feed)
    with open(mapping, "w", encoding="utf-8") as f:
        f.write("sku\tcustomization_template\tfields\n")
        for c in custom_rows:
            f.write(f"{c['sku']}\t{c['template']}\t{c['fields']}\n")
    summary = {"products_in": len(products), "amazon_rows": len(listings),
               "parents": sum(r["parentage"] == "parent" for r in listings),
               "personalised_skus": len(custom_rows), "custom_templates": list(templates),
               "held_back": len(held), "checks": len(checks), "files": [xl, feed, mapping]}
    if a.template:
        outp = os.path.join(a.out, "Amazon template FILLED - " + os.path.basename(a.template))
        filled, empty = fill_template(a.template, listings, cfg, outp)
        summary["files"].append(outp)
        summary["template_columns_filled"] = len(filled)
        summary["template_columns_left_blank"] = empty[:60]
    json.dump(summary, open(os.path.join(a.out, "summary.json"), "w"), indent=1)
    print(json.dumps(summary, indent=1))


if __name__ == "__main__":
    main()
