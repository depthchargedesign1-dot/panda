"""Build Amazon UK upload files (listings + Amazon Custom spec) from Shopify products.

What it makes, in the output folder:
  amazon-upload.xlsx            Listings (parent + child rows, Amazon flat-file column names), Customisation,
                                Column guide and ASK sheets. Cells still marked ASK are filled yellow.
  amazon-upload.csv             The Listings sheet as CSV (for reading / copying; Amazon itself only takes
                                its own .xlsm/.xlsx template or tab-delimited .txt).
  amazon-customisation-spec.csv Every Amazon Custom field per parent, mapped from foxy.personalise_fields
                                (same text/textarea/image rules and character limits as the live theme).
  build-report.txt              Title / bullet / search-term checks and anything left as ASK.

Usage
  # 1. Get the Shopify data. Either run QUERY (print it with --print-query) through the Shopify MCP
  #    graphql_query tool with {"ids": [...]} and save the raw response as JSON, or let the script fetch it:
  #    SHOPIFY_SHOP=xxx.myshopify.com SHOPIFY_ADMIN_TOKEN=shpat_... python3 tools/amazon_export.py --fetch ID ID ...
  # 2. Build:
  python3 tools/amazon_export.py --source exports/amazon/<date>/shopify-source.json \
      --config exports/amazon/<date>/amazon-config.json --out exports/amazon/<date> --name amazon-test-upload

Config (optional JSON). Without it every product becomes its own parent, variations come from its Shopify
options. With it you can group several Shopify products into one Amazon parent (e.g. 8 hat colours that are
separate products on Shopify) and set Amazon-only fields:
  {"defaults": {"brand_name": "Foxy Printing", "quantity": 20, "fulfillment_latency": "ASK", ...},
   "groups": [{"key": "hat", "ids": ["gid://shopify/Product/1", ...], "feed_product_type": "HAT",
               "variation_theme": "Color", "child_values": {"gid://...": "Black & Gold"},
               "title": "...", "child_title": "{parent} - {value}", "title_max": 125,
               "extra_keywords": "...", "attributes": {"material_type": "Acrylic"}, "weight_g": 85}]}

Rules built in (CLAUDE.md + Amazon listing rules, see exports/amazon/2026-10-06-test/README.md):
  - item_sku = Shopify SKU (optional --sku-prefix, e.g. AMZ-, only if they clash with SKUs already on Amazon).
  - Titles: brand first, no promo words, no ! $ ? _ { } ^ ¬ ¦, no word more than twice, <= 200 (125 apparel).
  - Shopify-only text is removed: live preview, basket/checkout, multi-buy, phone numbers, "contact us",
    delivery promises and cross-sell lines. Personalisation text is rewritten for Amazon's "Customise now".
  - Third-party-name products keep their disclaimer (<p class="disclaimer">) at the end of the description.
  - Nothing is invented: unknown values are written as ASK.
"""
import argparse, csv, html, json, os, re, sys, urllib.request

QUERY = """query($ids: [ID!]!) { nodes(ids: $ids) { ... on Product {
  id handle title status productType vendor tags templateSuffix descriptionHtml seo { title description }
  options { name values }
  media(first: 20) { nodes { ... on MediaImage { image { url altText } } } }
  variants(first: 100) { nodes { sku title price selectedOptions { name value } image { url }
    inventoryItem { measurement { weight { unit value } } } } }
  foxy: metafields(first: 20, namespace: "foxy") { nodes { key value } }
  g: metafields(first: 20, namespace: "mm-google-shopping") { nodes { key value } } } } }"""

ASK = "ASK"
DEFAULTS = {
    "brand_name": "Foxy Printing", "manufacturer": "Foxy Printing", "quantity": 20,
    "fulfillment_latency": ASK, "merchant_shipping_group_name": ASK, "recommended_browse_nodes": ASK,
    "country_of_origin": ASK, "condition_type": "New", "currency": "GBP", "update_delete": "Update",
}
# Best guesses at Amazon product type names - confirm in Seller Central's template generator.
PRODUCT_TYPE_GUESS = {"mugs": "DRINKING_CUP", "terrace flags": "FLAG", "flags": "FLAG", "bobble hats": "HAT",
                      "hats": "HAT", "caps": "HAT", "face masks": "COSTUME_MASK", "celebrity masks": "COSTUME_MASK"}
THEME_COLUMN = {"color": "color_name", "colour": "color_name", "size": "size_name", "style": "style_name"}

COLUMNS = [
    "feed_product_type", "item_sku", "update_delete", "brand_name", "manufacturer", "part_number",
    "external_product_id", "external_product_id_type", "item_name",
    "parent_child", "parent_sku", "relationship_type", "variation_theme",
    "color_name", "color_map", "size_name", "style_name",
    "standard_price", "currency", "quantity", "fulfillment_latency", "condition_type",
    "merchant_shipping_group_name", "recommended_browse_nodes",
    "product_description", "bullet_point1", "bullet_point2", "bullet_point3", "bullet_point4", "bullet_point5",
    "generic_keywords", "main_image_url", "swatch_image_url",
] + [f"other_image_url{i}" for i in range(1, 9)] + [
    "item_weight", "item_weight_unit_of_measure", "country_of_origin",
]
EXTRA_COLUMNS_ORDER = ["material_type", "capacity", "care_instructions", "target_gender", "age_range_description",
                       "department_name"]

COLUMN_GUIDE = {
    "feed_product_type": ("Product Type", "Amazon product type. Best guess - confirm when you generate the template."),
    "item_sku": ("SKU", "Shopify SKU. Parent SKU = child SKU without the last -NN."),
    "update_delete": ("Listing Action", "Update = create or full update. ('Create or Replace (Full Update)' in new templates.)"),
    "brand_name": ("Brand Name", "Must match Brand Registry / GTIN exemption brand exactly. ASK if registered."),
    "manufacturer": ("Manufacturer", ""),
    "part_number": ("Part Number / Manufacturer Part Number", "Same as SKU (we have no other MPN)."),
    "external_product_id": ("External Product ID", "Left blank: no barcodes. Needs a GTIN exemption first."),
    "external_product_id_type": ("External Product ID Type", "Blank (GTIN exempt)."),
    "item_name": ("Item Name (Title)", "Checked: <=200 chars (125 for hats), no promo words or banned symbols."),
    "parent_child": ("Parentage Level", "parent / child"),
    "parent_sku": ("Parent SKU", "Children only."),
    "relationship_type": ("Child Relationship Type", "Variation (children only)."),
    "variation_theme": ("Variation Theme Name", "Color / Size / Style. Pick the matching value from the template dropdown."),
    "color_name": ("Colour", ""), "color_map": ("Colour Map", "Standard colour family."),
    "size_name": ("Size", ""), "style_name": ("Style", "Used for the mug's country band (Amazon has no 'Country' theme)."),
    "standard_price": ("Your Price GBP (Sell on Amazon, UK)", "Shopify price. Children only."),
    "currency": ("Currency", ""),
    "quantity": ("Quantity (UK)", "Made to order - a working stock figure."),
    "fulfillment_latency": ("Handling Time (UK)", "ASK - days from order to dispatch."),
    "condition_type": ("Offering Condition Type", "New"),
    "merchant_shipping_group_name": ("Merchant Shipping Group (UK)", "ASK - name of your shipping template."),
    "recommended_browse_nodes": ("Recommended Browse Nodes", "ASK - from the template's browse tree / product type picker."),
    "product_description": ("Product Description", "Plain text, <=2000 characters."),
    "bullet_point1": ("Bullet Point", "Five bullets, <=250 characters each."),
    "generic_keywords": ("Generic Keyword", "Backend search terms, <=249 bytes, words not already in the title where possible."),
    "main_image_url": ("Main Image URL", "Shopify CDN. Main image must be on pure white, no text."),
    "swatch_image_url": ("Swatch Image URL", "Colour variations only."),
    "other_image_url1": ("Other Image URL", "Up to 8."),
    "item_weight": ("Item Weight", "ASK where Shopify has no real weight."),
    "country_of_origin": ("Country/Region of Origin", "ASK (where the blank was made)."),
}

PROMO_WORDS = re.compile(r"\b(best[- ]?sell\w*|best|cheap\w*|free (?:shipping|delivery|p&p)|sale|top[- ]rated|hot|"
                         r"offer|deal|discount|official|licen[cs]ed|authentic|genuine|approved|endorsed|merchandise|"
                         r"fast turnaround|guarantee\w*)\b", re.I)
BANNED_CHARS = set("!$?_{}^¬¦")
SMALL_WORDS = {"a", "an", "the", "and", "or", "for", "with", "of", "to", "in", "on", "by", "&", "-", "x"}
# Sentences / clauses that only make sense on our own website.
DROP = re.compile(r"live preview|basket|checkout|multi-?buy|\bsave \d+%|\b0\d{3,4}\s?\d{3}\s?\d{3,4}\b|\bring us\b|"
                  r"\bring \d|give us a ring|contact us|our (?:football )?range|our website|first box|message box|"
                  r"delivery is free|free uk delivery|postage|royal mail|fast turnaround|post it out", re.I)
SKIP_TAG = re.compile(r"^(range-|io-|machine-|country-|foxy-|dcd |clothing \d{4}|third-party-name)|/|"
                      r"^[a-z]+-[a-z-]+$", re.I)


# ---------------------------------------------------------------- Shopify data
def fetch(ids):
    shop, token = os.environ.get("SHOPIFY_SHOP"), os.environ.get("SHOPIFY_ADMIN_TOKEN")
    if not (shop and token):
        sys.exit("Set SHOPIFY_SHOP and SHOPIFY_ADMIN_TOKEN, or run --print-query through the Shopify MCP and use --source.")
    ids = [i if i.startswith("gid://") else f"gid://shopify/Product/{i}" for i in ids]
    req = urllib.request.Request(f"https://{shop}/admin/api/2025-07/graphql.json", method="POST",
                                 data=json.dumps({"query": QUERY, "variables": {"ids": ids}}).encode(),
                                 headers={"Content-Type": "application/json", "X-Shopify-Access-Token": token})
    return json.load(urllib.request.urlopen(req))


def load_products(raw):
    nodes = raw.get("data", raw).get("nodes", raw.get("nodes", []))
    out = {}
    for p in nodes:
        if not p:
            continue
        p["images"] = [m["image"] for m in (p.get("media") or {}).get("nodes", []) if m and m.get("image")]
        p["variants_"] = (p.get("variants") or {}).get("nodes", [])
        p["foxy_"] = {m["key"]: m["value"] for m in (p.get("foxy") or {}).get("nodes", [])}
        p["google_"] = {m["key"]: m["value"] for m in (p.get("g") or {}).get("nodes", [])}
        out[p["id"]] = p
    return out


# ---------------------------------------------------------------- personalisation fields (same rules as theme)
def field_spec(label, index):
    """Mirror of sections/main-product.liquid: kind, max length, required (first field only)."""
    l = label.lower()
    words = re.findall(r"[a-z]+", l)
    kind = "text"
    if any(k in l for k in ("photo", "upload", "logo", "artwork", "badge", "crest")):
        kind = "image"
    elif any(k in l for k in ("message", "list", "reasons", "achievements", "notes", "how many")):
        kind = "textarea"
    mx = 24
    if kind == "textarea":
        mx = 160
    elif "age" in words:
        mx = 3
    elif "number" in words and "phone" not in l:
        mx = 3
    elif "initials" in words:
        mx = 4
    elif any(k in l for k in ("slogan", "team name", "colours")):
        mx = 40
    required = index == 0 and "optional" not in l
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", label)
    short, hint = (m.group(1), m.group(2)) if m else (label, "")
    hint = "" if hint.lower() == "optional" else re.sub(r",?\s*optional", "", hint, flags=re.I)
    return {"label": label, "short": short.strip(), "hint": hint.strip(), "kind": kind,
            "max": None if kind == "image" else mx, "required": required}


def personalise_fields(p):
    raw = p["foxy_"].get("personalise_fields")
    if not raw:
        return []
    try:
        labels = json.loads(raw)
    except ValueError:
        labels = [raw]
    return [field_spec(l, i) for i, l in enumerate(labels)]


def customise_sentence(fields):
    parts = []
    for f in fields:
        s = f["short"][0].lower() + f["short"][1:]
        if f["kind"] == "image":
            noun = re.sub(r"\s*upload\s*", " ", s).strip()
            txt = f"upload your {noun}"
        elif "message" in s:
            txt = "leave us a message"
        else:
            txt = f"type the {s}"
        if not f["required"]:
            txt += " (optional)"
        parts.append(txt)
    if not parts:
        return ""
    lst = parts[0] if len(parts) == 1 else ", ".join(parts[:-1]) + " and " + parts[-1]
    return f"Click Customise now to {lst}."


# ---------------------------------------------------------------- text helpers
def clean(t):
    t = html.unescape(re.sub(r"<[^>]+>", "", t or ""))
    t = t.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    t = t.replace("–", "-").replace("—", "-").replace(" ", " ")
    return re.sub(r"\s+", " ", t).strip()


def sentences(t):
    return [s for s in re.split(r"(?<=[.!?])\s+", t) if s]


def strip_shop_text(t):
    """Remove sentences, then clauses (split on ; and dashes) that only make sense on our website."""
    keep = []
    for s in sentences(t):
        if not DROP.search(s):
            keep.append(s)
            continue
        clauses = [c for c in re.split(r";\s*|\s+-\s+", s.rstrip(".")) if c and not DROP.search(c)]
        if clauses and len(clauses[0].split()) > 3:
            keep.append("; ".join(clauses).rstrip(".") + ".")
    return " ".join(keep)


def parse_description(desc):
    """Split the Shopify description into intro / personal / why / details / disclaimer."""
    blocks = re.findall(r"<(h2|h3|p|li)([^>]*)>(.*?)</\1>", desc or "", re.S)
    out = {"intro": [], "personal": [], "why": [], "details": [], "disclaimer": [], "other": []}
    section, after_h2 = "intro", False
    for tag, attrs, inner in blocks:
        text = clean(inner)
        if tag == "h2":
            section, after_h2 = "h2", True
            continue
        if tag == "h3":
            h = text.lower()
            section = ("why" if "love" in h else "details" if "size" in h or "detail" in h else
                       "delivery" if "deliver" in h else "disclaimer" if "note" in h else "other")
            continue
        if "disclaimer" in attrs:
            out["disclaimer"].append(text)
        elif section == "intro" and tag == "p":
            out["intro"].append(text)
        elif section == "h2" and tag == "p" and after_h2:
            out["personal"].append(text)
            after_h2 = False
        elif section in ("why", "details") and tag == "li":
            out[section].append(text)
        elif section == "details" and tag == "p":
            out["other"].append(text)
    return out


def cap(s):
    return s[:1].upper() + s[1:] if s else s


def fit(items, limit, sep="; ", prefix=""):
    out = prefix
    for it in items:
        cand = out + (sep if out != prefix else "") + it
        if len(cand) > limit:
            break
        out = cand
    return out


# ---------------------------------------------------------------- titles & checks
def make_title(shopify_title, brand):
    t = clean(shopify_title)
    t = re.sub(r"\s*-\s*(fast turnaround|\d+ sizes|\d+ colours)\b", "", t, flags=re.I)
    if not t.lower().startswith(brand.lower()):
        t = f"{brand} {t}"
    return t


def check_title(t, limit):
    probs = []
    if len(t) > limit:
        probs.append(f"{len(t)} chars (limit {limit})")
    bad = sorted(set(c for c in t if c in BANNED_CHARS))
    if bad:
        probs.append("banned characters " + "".join(bad))
    m = PROMO_WORDS.search(t)
    if m:
        probs.append(f"promo/claim word '{m.group(0)}'")
    counts = {}
    for w in re.findall(r"[a-z0-9']+", t.lower()):
        if w not in SMALL_WORDS:
            counts[w] = counts.get(w, 0) + 1
    rep = [w for w, n in counts.items() if n > 2]
    if rep:
        probs.append("word used more than twice: " + ", ".join(rep))
    return probs


def search_terms(products, title, extra):
    title_words = set(re.findall(r"[a-z0-9]+", title.lower()))
    words, seen = [], set()
    tags = []
    for p in products:
        tags += [t for t in p.get("tags", []) if not SKIP_TAG.search(t)]
    for chunk in [extra or ""] + tags:
        for w in re.findall(r"[a-z0-9']+", chunk.lower()):
            if w in seen or w in title_words or w in SMALL_WORDS or len(w) < 3:
                continue
            seen.add(w)
            words.append(w)
    out = ""
    for w in words:
        if len((out + " " + w).strip().encode()) > 249:
            break
        out = (out + " " + w).strip()
    return out


# ---------------------------------------------------------------- build
def option_label(name):
    n = name.lower()
    return "Color" if n in ("colour", "color") else "Size" if n == "size" else "Style"


def weight_of(variant, override):
    if override:
        return override, "GR"
    w = ((variant.get("inventoryItem") or {}).get("measurement") or {}).get("weight") or {}
    val, unit = w.get("value") or 0, (w.get("unit") or "GRAMS").upper()
    grams = {"GRAMS": 1, "KILOGRAMS": 1000, "OUNCES": 28.35, "POUNDS": 453.6}.get(unit, 1) * float(val)
    return (round(grams), "GR") if grams >= 1 else (ASK, "GR")


def build_content(p, fields, defaults):
    d = parse_description(p.get("descriptionHtml"))
    intro = " ".join(strip_shop_text(x) for x in d["intro"])
    keep_personal = [s for x in d["personal"] for s in sentences(x)
                     if re.search(r"right to use|print exactly", s, re.I)]
    if fields:
        personal = customise_sentence(fields)
        if not any("print exactly" in s.lower() for s in keep_personal):
            what = "upload" if all(f["kind"] == "image" for f in fields if f["required"]) else "type"
            keep_personal.insert(0, f"We print exactly what you {what}, so please check names and spelling before you order.")
        personal = " ".join([personal] + [s.replace("before you order", "before you order") for s in keep_personal])
    else:
        personal = " ".join(strip_shop_text(x) for x in d["personal"])
    why = [cap(strip_shop_text(x).rstrip(".")) for x in d["why"]]
    why = [w for w in why if w and not DROP.search(w)]
    details = [strip_shop_text(x).rstrip(".") for x in d["details"]]
    details = [x for x in details if x]
    disclaimer = " ".join(d["disclaimer"])

    bullets = []
    if fields:
        first = personal.split('. ')[0].rstrip('.')
        bullets.append(fit([f"Personalised for you: {first[:1].lower() + first[1:]}"], 250))
    for w in why:
        if len(bullets) >= 4:
            break
        bullets.append(w[:250])
    bullets.append(fit(details, 250, prefix="Size and details: "))
    while len(bullets) < 5 and len(why) > len(bullets) - 1:
        bullets.insert(-1, why[len(bullets) - 1][:250])
    bullets = (bullets + [""] * 5)[:5]

    paras = [intro, personal]
    if details:
        paras.append("Size and details: " + "; ".join(details) + ".")
    paras.append("Each one is made to order, so please allow for the dispatch time shown.")
    if disclaimer:
        paras.append("Please note: " + disclaimer)
    desc = "\n\n".join(x for x in paras if x)
    if len(desc) > 2000:
        desc = desc[:1990].rsplit(". ", 1)[0] + "."
    return desc, bullets


def build(products, config, sku_prefix=""):
    defaults = dict(DEFAULTS, **(config.get("defaults") or {}))
    groups = config.get("groups") or [{"ids": [pid]} for pid in products]
    rows, spec, report, extra_cols = [], [], [], []
    for g in groups:
        ps = [products[i] for i in g["ids"] if i in products]
        missing = [i for i in g["ids"] if i not in products]
        if missing:
            report.append(f"MISSING from source: {missing}")
        if not ps:
            continue
        lead = ps[0]
        brand = defaults["brand_name"]
        ptype = g.get("feed_product_type") or PRODUCT_TYPE_GUESS.get(lead.get("productType", "").lower(), ASK)
        title_max = g.get("title_max", 200)
        attrs = g.get("attributes") or {}
        for k in attrs:
            if k not in extra_cols:
                extra_cols.append(k)

        # children: one per Shopify variant, or one per product when several products are grouped
        children = []
        if len(ps) > 1:
            theme = g.get("variation_theme", "Color")
            for p in ps:
                v = p["variants_"][0]
                children.append((p, v, (g.get("child_values") or {}).get(p["id"]) or p["google_"].get("color", ASK)))
        else:
            opts = [o for o in lead.get("options", []) if o["name"] != "Title"]
            theme = g.get("variation_theme") or (option_label(opts[0]["name"]) if len(opts) == 1 else
                                                 "-".join(option_label(o["name"]) for o in opts) if opts else "")
            for v in lead["variants_"]:
                val = " / ".join(o["value"] for o in v.get("selectedOptions", []) if o["name"] != "Title")
                children.append((lead, v, val))
        single = len(children) == 1 and not theme

        parent_sku = sku_prefix + re.sub(r"-\d+$", "", children[0][1]["sku"])
        parent_title = g.get("title") or make_title(lead["title"], brand)
        fields = personalise_fields(lead)
        fmt = g.get("child_title", "{parent} - {value}")
        theme_col = THEME_COLUMN.get(theme.lower(), "style_name") if theme else None

        def base_row(p, sku, title):
            desc, bullets = build_content(p, fields, defaults)
            kw = search_terms(ps, title, g.get("extra_keywords"))
            r = {c: "" for c in COLUMNS}
            r.update({"feed_product_type": ptype, "item_sku": sku, "update_delete": defaults["update_delete"],
                      "brand_name": brand, "manufacturer": defaults["manufacturer"], "item_name": title,
                      "product_description": desc, "generic_keywords": kw})
            for i, b in enumerate(bullets, 1):
                r[f"bullet_point{i}"] = b
            r.update(attrs)
            return r

        def images(p, first=None):
            urls = [i["url"] for i in p["images"]]
            if first:
                urls = [first] + [u for u in urls if u != first]
            return urls

        if not single:
            pr = base_row(lead, parent_sku, parent_title)
            pr.update({"parent_child": "parent", "variation_theme": theme})
            imgs = images(lead) if len(ps) == 1 else [c[0]["images"][0]["url"] for c in children if c[0]["images"]]
            if imgs:
                pr["main_image_url"] = imgs[0]
                for i, u in enumerate(imgs[1:9], 1):
                    pr[f"other_image_url{i}"] = u
            rows.append(pr)
            for prob in check_title(parent_title, title_max):
                report.append(f"{parent_sku} title: {prob}")

        for p, v, val in children:
            sku = sku_prefix + v["sku"]
            title = parent_title if single else (
                make_title(p["title"], brand) if len(ps) > 1 and not g.get("child_title") else
                fmt.format(parent=parent_title, value=val))
            r = base_row(p, sku, title)
            vimg = (v.get("image") or {}).get("url")
            imgs = images(p, vimg)
            if imgs:
                r["main_image_url"] = imgs[0]
                for i, u in enumerate(imgs[1:9], 1):
                    r[f"other_image_url{i}"] = u
            wt, unit = weight_of(v, g.get("weight_g"))
            r.update({"part_number": sku, "standard_price": v.get("price", ""), "currency": defaults["currency"],
                      "quantity": defaults["quantity"], "fulfillment_latency": defaults["fulfillment_latency"],
                      "condition_type": defaults["condition_type"],
                      "merchant_shipping_group_name": defaults["merchant_shipping_group_name"],
                      "recommended_browse_nodes": g.get("recommended_browse_nodes", defaults["recommended_browse_nodes"]),
                      "item_weight": wt, "item_weight_unit_of_measure": unit,
                      "country_of_origin": g.get("country_of_origin", defaults["country_of_origin"])})
            if not single:
                r.update({"parent_child": "child", "parent_sku": parent_sku, "relationship_type": "Variation",
                          "variation_theme": theme, theme_col: val})
                if theme_col == "color_name":
                    r["color_map"] = (g.get("color_map") or {}).get(p["id"]) or p["google_"].get("color", "")
                    r["swatch_image_url"] = r["main_image_url"]
            for prob in check_title(title, title_max):
                report.append(f"{sku} title: {prob}")
            for i in range(1, 6):
                if len(r[f"bullet_point{i}"]) > 250:
                    report.append(f"{sku} bullet {i} over 250 chars")
            if "third-party-name" in p.get("tags", []) and not p["descriptionHtml"].count("disclaimer"):
                report.append(f"{sku}: tagged third-party-name but no disclaimer in Shopify description")
            rows.append(r)

        # Amazon Custom spec
        child_skus = " ".join(sku_prefix + c[1]["sku"] for c in children)
        if fields:
            for n, f in enumerate(fields, 1):
                spec.append({
                    "amazon_parent_sku": parent_sku if not single else child_skus, "apply_to_child_skus": child_skus,
                    "shopify_product": lead["handle"] + (f" (+{len(ps) - 1} more)" if len(ps) > 1 else ""),
                    "field_order": n,
                    "amazon_custom_type": {"image": "Image (customer upload)", "textarea": "Text (multi-line)",
                                           "text": "Text (single line)"}[f["kind"]],
                    "label": f["short"], "instructions_placeholder": f["hint"],
                    "max_characters": f["max"] or "", "max_lines": (4 if f["kind"] == "textarea" else 1) if f["max"] else "",
                    "required": "Yes" if f["required"] else "No", "price_surcharge_gbp": "0.00",
                    "source": f"foxy.personalise_fields: {f['label']}",
                    "notes": ("Set the print area (surface + placement box) on the product image in Seller Central."
                              if f["kind"] == "image" else ""),
                })
        else:
            spec.append({"amazon_parent_sku": parent_sku, "apply_to_child_skus": child_skus,
                         "shopify_product": lead["handle"], "field_order": "", "amazon_custom_type": "None",
                         "label": "", "instructions_placeholder": "", "max_characters": "", "max_lines": "",
                         "required": "", "price_surcharge_gbp": "",
                         "source": "no foxy.personalise_fields",
                         "notes": "Not a personalised product - do not enable Amazon Custom."})
        if theme:
            vals = " | ".join(c[2] for c in children)
            spec.append({"amazon_parent_sku": parent_sku, "apply_to_child_skus": child_skus,
                         "shopify_product": lead["handle"], "field_order": "variation",
                         "amazon_custom_type": f"Variation ({theme}) - in the upload file, not Amazon Custom",
                         "label": theme, "instructions_placeholder": vals, "max_characters": "", "max_lines": "",
                         "required": "Yes", "price_surcharge_gbp": "0.00 (each child has its own price)",
                         "source": "Shopify option" + ("s" if len(ps) == 1 else ": separate Shopify products"),
                         "notes": "Customers pick this on the Amazon page before Customise now."})
    asks = {}
    for r in rows:
        for k, v in r.items():
            if v == ASK:
                asks.setdefault(k, []).append(r["item_sku"])
    for k, skus in asks.items():
        report.append(f"ASK {k}: {len(skus)} rows ({', '.join(skus)})")
    return rows, spec, report, extra_cols


SPEC_COLS = ["amazon_parent_sku", "apply_to_child_skus", "shopify_product", "field_order", "amazon_custom_type",
             "label", "instructions_placeholder", "max_characters", "max_lines", "required", "price_surcharge_gbp",
             "source", "notes"]


def write(out, name, rows, spec, report, extra_cols):
    os.makedirs(out, exist_ok=True)
    cols = COLUMNS + [c for c in EXTRA_COLUMNS_ORDER if c in extra_cols] + \
        [c for c in extra_cols if c not in EXTRA_COLUMNS_ORDER]
    with open(os.path.join(out, f"{name}.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    with open(os.path.join(out, "amazon-customisation-spec.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, SPEC_COLS)
        w.writeheader()
        w.writerows(spec)
    with open(os.path.join(out, "build-report.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(report or ["No problems found."]) + "\n")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("openpyxl not installed - CSV only (pip install openpyxl)")
        return cols
    wb = Workbook()
    yellow = PatternFill("solid", fgColor="FFF2A8")
    head = PatternFill("solid", fgColor="F2D7C9")

    def sheet(ws, header, data):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), head
        for r in data:
            ws.append([r.get(h, "") for h in header])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.value == ASK:
                    c.fill = yellow
                c.alignment = Alignment(wrap_text=False, vertical="top")
        ws.freeze_panes = "C2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or ""))
                                                                                   for c in col[:3]) + 2))
    ws = wb.active
    ws.title = "Listings"
    sheet(ws, cols, rows)
    sheet(wb.create_sheet("Customisation"), SPEC_COLS, spec)
    g = wb.create_sheet("Column guide")
    sheet(g, ["our column (flat-file name)", "label in Amazon's 2025-26 template (check)", "notes"],
          [{"our column (flat-file name)": c, "label in Amazon's 2025-26 template (check)": COLUMN_GUIDE.get(
              re.sub(r"\d+$", "1", c) if c.startswith(("bullet", "other_image")) else c, ("", ""))[0],
            "notes": COLUMN_GUIDE.get(re.sub(r"\d+$", "1", c) if c.startswith(("bullet", "other_image")) else c,
                                      ("", ""))[1]} for c in cols])
    a = wb.create_sheet("ASK")
    sheet(a, ["item", "still to fill"], [{"item": x.split(":")[0], "still to fill": x.split(":", 1)[-1].strip()}
                                         for x in report if x.startswith("ASK")])
    wb.save(os.path.join(out, f"{name}.xlsx"))
    return cols


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", help="saved GraphQL response (QUERY) as JSON")
    ap.add_argument("--fetch", nargs="*", help="Shopify product IDs to fetch with SHOPIFY_SHOP/SHOPIFY_ADMIN_TOKEN")
    ap.add_argument("--config", help="optional JSON config (groups, defaults)")
    ap.add_argument("--out", default="exports/amazon/out")
    ap.add_argument("--name", default="amazon-upload")
    ap.add_argument("--sku-prefix", default="", help="e.g. AMZ- if Shopify SKUs clash with existing Amazon SKUs")
    ap.add_argument("--print-query", action="store_true")
    a = ap.parse_args()
    if a.print_query:
        print(QUERY)
        return
    config = json.load(open(a.config)) if a.config else {}
    if a.fetch:
        raw = fetch(a.fetch)
        os.makedirs(a.out, exist_ok=True)
        json.dump(raw, open(os.path.join(a.out, "shopify-source.json"), "w"), ensure_ascii=False, indent=1)
    elif a.source:
        raw = json.load(open(a.source))
    else:
        ap.error("give --source or --fetch")
    products = load_products(raw)
    rows, spec, report, extra = build(products, config, a.sku_prefix)
    write(a.out, a.name, rows, spec, report, extra)
    print(f"{len(rows)} listing rows, {len(spec)} customisation rows -> {a.out}")
    print(f"{len(report)} report lines (see build-report.txt)")


if __name__ == "__main__":
    main()
